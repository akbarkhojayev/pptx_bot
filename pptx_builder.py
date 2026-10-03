"""Prezentatsiya (.pptx) yasovchi modul.

Kirish ma'lumoti -- AI yoki Wikipedia'dan olingan "deck" lug'ati:
    {"title", "subtitle", "palette", "slides": [{"type": ..., "title": ..., "notes": ..., ...}]}
Har bir slayd turi (bullets, cards, stats, timeline, compare, quote, conclusion) o'z maketiga ega,
shuning uchun natija bir xil "matn devori" emas, balki haqiqiy taqdimotga o'xshaydi.
"""
import logging
import math
import os
import re
import textwrap
from datetime import date
from typing import List, Optional, Sequence, Tuple

from lxml import etree
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

# ---------------- PALETTES ----------------
# AI mavzuga qarab palitra nomini tanlaydi; har biri: dark (sarlavha fon), primary (asosiy rang),
# accent (bezak), tint (kartochka foni), head (sarlavha shrifti).
PALETTES = {
    "ocean":      {"dark": "0B2545", "primary": "13507C", "accent": "1FA2D8", "tint": "EAF3F9", "head": "Calibri"},
    "forest":     {"dark": "1B3A1C", "primary": "2C5F2D", "accent": "97BC62", "tint": "EEF4E8", "head": "Cambria"},
    "terracotta": {"dark": "4A1F17", "primary": "A3452F", "accent": "E0A458", "tint": "F8EEEA", "head": "Cambria"},
    "royal":      {"dark": "21114A", "primary": "4B2A8C", "accent": "B98CF0", "tint": "F1ECFA", "head": "Calibri"},
    "midnight":   {"dark": "131A33", "primary": "1E2761", "accent": "F2A541", "tint": "ECEEF7", "head": "Calibri"},
    "teal":       {"dark": "0B3B40", "primary": "02707D", "accent": "02C39A", "tint": "E5F4F3", "head": "Calibri"},
    "berry":      {"dark": "3A1626", "primary": "6D2E46", "accent": "D8A48F", "tint": "F6EDF0", "head": "Cambria"},
    "cherry":     {"dark": "4D0009", "primary": "990011", "accent": "E8B04B", "tint": "FBEDEE", "head": "Cambria"},
    "charcoal":   {"dark": "1F262B", "primary": "36454F", "accent": "F25C54", "tint": "EEF1F3", "head": "Arial"},
}
DEFAULT_PALETTE = "ocean"

TEXT = "1D2433"
MUTED = "6B7280"
WHITE = "FFFFFF"
FONT_BODY = "Calibri"
LANG = "uz-Latn-UZ"

# ---------------- GEOMETRY (16:9, 13.33" x 7.5") ----------------
SW = Inches(13.3333)
SH = Inches(7.5)
MX = Inches(0.6)                 # chap/o'ng chekka
CW = SW - 2 * MX                 # kontent kengligi
TITLE_Y = Inches(0.62)
TITLE_H = Inches(1.0)
CY = Inches(1.85)                # kontent boshlanishi
CB = Inches(6.8)                 # kontent oxiri
CH = CB - CY
GAP = Inches(0.3)

BODY_SIZES = (24, 22, 20, 19, 18, 17, 16, 15, 14)
SMALL_SIZES = (17, 16, 15, 14, 13)


def rgb(hex_str: str) -> RGBColor:
    return RGBColor.from_string(hex_str)


def mix(hex_a: str, hex_b: str, t: float) -> str:
    """hex_a dan hex_b ga t (0..1) nisbatda aralashtirilgan rang."""
    a = [int(hex_a[i:i + 2], 16) for i in (0, 2, 4)]
    b = [int(hex_b[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(x + (y - x) * t):02X}" for x, y in zip(a, b))


# ---------------- TEXT FITTING ----------------
def _count_lines(text: str, chars_per_line: float) -> int:
    width = max(1, int(chars_per_line))
    return sum(max(1, len(textwrap.wrap(line, width=width))) for line in text.split("\n"))


def fit_size(
    paragraphs: Sequence[str],
    box_w: int,
    box_h: int,
    sizes: Sequence[int],
    line_spacing: float = 1.1,
    space_after_ratio: float = 0.5,
    char_w: float = 0.45,
    indent: int = 0,
) -> int:
    """Matn qutiga sig'adigan eng katta shrift o'lchamini (pt) tanlaydi.
    O'rtacha belgi kengligi ~0.5x shrift, qator balandligi ~1.2x shrift x line_spacing
    (LibreOffice'da o'lchangan qiymatlarga moslangan)."""
    for size in sizes:
        if text_height(paragraphs, box_w, size, line_spacing, space_after_ratio, char_w, indent) <= box_h * 0.95:
            return size
    return sizes[-1]


def text_height(paragraphs: Sequence[str], box_w: int, size: float, line_spacing: float = 1.1,
                space_after_ratio: float = 0.5, char_w: float = 0.45, indent: int = 0) -> int:
    """Berilgan shriftda matn egallaydigan taxminiy balandlik (EMU)."""
    w_pt = (box_w - indent) / 12700
    lines = sum(_count_lines(p, w_pt / (size * char_w)) for p in paragraphs)
    needed_pt = lines * size * 1.2 * line_spacing + max(0, len(paragraphs) - 1) * size * space_after_ratio
    return int(needed_pt * 12700)


def fit_single_line(text: str, box_w: int, sizes: Sequence[int], char_w: float = 0.56) -> int:
    w_pt = box_w / 12700
    for size in sizes:
        if len(text) * size * char_w <= w_pt:
            return size
    return sizes[-1]


# ---------------- LOW-LEVEL SHAPE HELPERS ----------------
def _style_run(run, size: float, color: str, font: str = FONT_BODY, bold: bool = False, italic: bool = False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.name = font
    f.color.rgb = rgb(color)
    run._r.get_or_add_rPr().set("lang", LANG)


def _prep_frame(tf, anchor=MSO_ANCHOR.TOP, autofit: bool = False, margin: int = 0):
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = margin
    tf.auto_size = MSO_AUTO_SIZE.NONE
    if autofit:
        # Zaxira: matn baribir sig'masa, PowerPoint/LibreOffice uni kichraytiradi.
        body_pr = tf._txBody.find(qn("a:bodyPr"))
        for child in list(body_pr):
            if child.tag in (qn("a:noAutofit"), qn("a:spAutoFit"), qn("a:normAutofit")):
                body_pr.remove(child)
        body_pr.append(OxmlElement("a:normAutofit"))


def _set_bullet(paragraph, char: str, color: str, indent: int):
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.set("marL", str(int(indent)))
    p_pr.set("indent", str(-int(indent)))
    bu_clr = OxmlElement("a:buClr")
    clr = OxmlElement("a:srgbClr")
    clr.set("val", color)
    bu_clr.append(clr)
    bu_font = OxmlElement("a:buFont")
    bu_font.set("typeface", "Arial")
    bu_char = OxmlElement("a:buChar")
    bu_char.set("char", char)
    for el in (bu_clr, bu_font, bu_char):
        p_pr.append(el)


def _strip_bullet(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buClr", "a:buFont", "a:buChar"):
        for el in p_pr.findall(qn(tag)):
            p_pr.remove(el)
    p_pr.set("marL", "0")
    p_pr.set("indent", "0")
    p_pr.append(OxmlElement("a:buNone"))


def _no_bullet(paragraph):
    paragraph._p.get_or_add_pPr().append(OxmlElement("a:buNone"))


def add_box(slide, x, y, w, h, fill: str, shape=MSO_SHAPE.RECTANGLE, alpha: Optional[int] = None,
            radius: Optional[float] = None, name: Optional[str] = None):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = rgb(fill)
    sp.line.fill.background()
    sp.shadow.inherit = False
    # Shablon uslubidagi soya (effectRef) ham o'chiriladi -- aks holda ba'zi dasturlar uni baribir chizadi.
    style = sp._element.find(qn("p:style"))
    if style is not None:
        style.find(qn("a:effectRef")).set("idx", "0")
    if alpha is not None:
        srgb = sp.fill.fore_color._xFill.find(qn("a:srgbClr"))
        a = OxmlElement("a:alpha")
        a.set("val", str(int(alpha * 1000)))
        srgb.append(a)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = radius
    if name:
        sp.name = name
    return sp


def add_text(slide, x, y, w, h, paragraphs: Sequence[str], *, size: float, color: str, font: str = FONT_BODY,
             bold: bool = False, italic: bool = False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             line_spacing: float = 1.1, space_after_ratio: float = 0.5, autofit: bool = False,
             name: Optional[str] = None):
    box = slide.shapes.add_textbox(x, y, w, h)
    if name:
        box.name = name
    tf = box.text_frame
    _prep_frame(tf, anchor, autofit)
    for i, text in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        if i < len(paragraphs) - 1:
            p.space_after = Pt(size * space_after_ratio)
        _style_run(p.add_run(), size, color, font, bold, italic)
        p.runs[0].text = text
    return box


def _write_bullets(tf, bullets: Sequence[str], size: int, color: str, bullet_color: str,
                   char: str = "•", lead_color: Optional[str] = None):
    """Har bir punkt -- alohida paragraf, haqiqiy PowerPoint bullet bilan.
    "Atama: izoh" ko'rinishidagi punktlarda atama qalin yoziladi."""
    indent = Pt(size * 1.1)
    for i, text in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.1
        if i < len(bullets) - 1:
            p.space_after = Pt(size * 0.6)
        m = re.match(r"^([^:]{2,40}):\s+(.+)$", text)
        if m:
            r1 = p.add_run()
            r1.text = m.group(1) + ": "
            _style_run(r1, size, lead_color or color, bold=True)
            r2 = p.add_run()
            r2.text = m.group(2)
            _style_run(r2, size, color)
        else:
            r = p.add_run()
            r.text = text
            _style_run(r, size, color)
        _set_bullet(p, char, bullet_color, indent)


def add_bullets(slide, x, y, w, h, bullets: Sequence[str], pal: dict, sizes=BODY_SIZES, char="•",
                color: str = TEXT, name: str = "Bullets") -> int:
    size = fit_size(bullets, w, h, sizes, line_spacing=1.1, space_after_ratio=0.6, indent=Pt(sizes[0] * 1.1))
    box = slide.shapes.add_textbox(x, y, w, h)
    box.name = name
    tf = box.text_frame
    _prep_frame(tf, autofit=True)
    _write_bullets(tf, bullets, size, color, pal["primary"], char, lead_color=pal["dark"])
    return size


def add_picture_cover(slide, path: str, x, y, w, h, rounded: bool = False, name: str = "Picture"):
    """Rasmni qutini to'liq qoplaydigan qilib (proporsiyani buzmasdan, ortiqchasini kesib) joylaydi."""
    with Image.open(path) as im:
        iw, ih = im.size
    pic = slide.shapes.add_picture(path, x, y, w, h)
    pic.name = name
    box_ratio, img_ratio = w / h, iw / ih
    if img_ratio > box_ratio:
        c = (1 - box_ratio / img_ratio) / 2
        pic.crop_left = pic.crop_right = c
    elif img_ratio < box_ratio:
        c = (1 - img_ratio / box_ratio) / 2
        pic.crop_top = pic.crop_bottom = c
    if rounded:
        geom = pic._element.spPr.find(qn("a:prstGeom"))
        geom.set("prst", "roundRect")
        av = geom.find(qn("a:avLst"))
        if av is None:
            av = OxmlElement("a:avLst")
            geom.append(av)
        gd = OxmlElement("a:gd")
        gd.set("name", "adj")
        gd.set("fmla", "val 6000")
        av.append(gd)
    return pic


def add_number_circle(slide, x, y, d, number, fill: str, text_color: str = WHITE):
    c = add_box(slide, x, y, d, d, fill, MSO_SHAPE.OVAL, name=f"Number {number}")
    tf = c.text_frame
    _prep_frame(tf, MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(number)
    _style_run(r, max(11, round(d / 12700 * 0.42)), text_color, bold=True)
    return c


def set_gradient_bg(slide, c1: str, c2: str, angle: float):
    fill = slide.background.fill
    fill.gradient()
    fill.gradient_stops[0].color.rgb = rgb(c1)
    fill.gradient_stops[1].color.rgb = rgb(c2)
    fill.gradient_angle = angle


def set_solid_bg(slide, color: str):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(color)


def add_transition(slide):
    """Har bir slaydga yumshoq "Fade" o'tish effekti."""
    el = etree.fromstring(
        '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="med">'
        '<p:fade/></p:transition>'
    )
    sld = slide._element
    anchor = sld.find(qn("p:clrMapOvr"))
    if anchor is None:
        anchor = sld.find(qn("p:cSld"))
    anchor.addnext(el)


def set_notes(slide, text: str):
    if text:
        slide.notes_slide.notes_text_frame.text = text


def set_title(slide, text: str, x, y, w, h, sizes: Sequence[int], color: str, font: str,
              align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, char_w: float = 0.55):
    """Sarlavhani maketdagi haqiqiy title placeholder'ga yozadi -- PowerPoint'ning
    "Outline" va navigatsiya panelida slayd nomlari to'g'ri ko'rinadi."""
    ph = slide.shapes.title
    ph.left, ph.top, ph.width, ph.height = x, y, w, h
    tf = ph.text_frame
    _prep_frame(tf, anchor)
    size = fit_size([text], w, h, sizes, line_spacing=0.95, space_after_ratio=0, char_w=char_w)
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = 0.95
    r = p.add_run()
    r.text = text
    _style_run(r, size, color, font, bold=True)
    return ph


def apply_theme(prs, pal_name: str, pal: dict):
    """Fayl ichidagi Office mavzusi (rang sxemasi va shriftlar)ni palitraga moslaydi, shunda
    foydalanuvchi PowerPoint'da yangi shakl/jadval qo'shsa ham ranglar mos bo'ladi."""
    theme_part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(theme_part.blob)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    scheme = root.find(".//a:clrScheme", ns)
    scheme.set("name", pal_name.title())
    colors = {
        "dk1": "000000", "lt1": WHITE, "dk2": pal["dark"], "lt2": pal["tint"],
        "accent1": pal["primary"], "accent2": pal["accent"], "accent3": mix(pal["primary"], WHITE, 0.35),
        "accent4": mix(pal["accent"], "000000", 0.25), "accent5": mix(pal["dark"], WHITE, 0.3),
        "accent6": MUTED, "hlink": pal["primary"], "folHlink": mix(pal["primary"], WHITE, 0.3),
    }
    for tag, val in colors.items():
        el = scheme.find(f"a:{tag}", ns)
        if el is None:
            continue
        for child in list(el):
            el.remove(child)
        etree.SubElement(el, f"{{{ns['a']}}}srgbClr").set("val", val)
    root.find(".//a:fontScheme/a:majorFont/a:latin", ns).set("typeface", pal["head"])
    root.find(".//a:fontScheme/a:minorFont/a:latin", ns).set("typeface", FONT_BODY)
    theme_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def normalize_image(src: str, dst: str, max_side: int = 2000) -> str:
    """Yuklangan rasmni (har qanday format, EXIF burilishi bilan) PowerPoint uchun xavfsiz JPEG'ga aylantiradi."""
    with Image.open(src) as im:
        im = ImageOps.exif_transpose(im)
        if im.mode in ("RGBA", "LA", "P"):
            im = im.convert("RGBA")
            bg = Image.new("RGB", im.size, (255, 255, 255))
            bg.paste(im, mask=im.split()[-1])
            im = bg
        else:
            im = im.convert("RGB")
        im.thumbnail((max_side, max_side))
        im.save(dst, "JPEG", quality=90)
    return dst


# ---------------- SLIDE BUILDER ----------------
class DeckBuilder:
    def __init__(self, deck: dict):
        self.deck = deck
        self.topic = deck["title"]
        self.pal_name = deck.get("palette") if deck.get("palette") in PALETTES else DEFAULT_PALETTE
        self.pal = PALETTES[self.pal_name]
        self.head = self.pal["head"]
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = SW, SH
        self.page = 0
        apply_theme(self.prs, self.pal_name, self.pal)

    # -------- common frames --------
    def _new(self, layout_idx: int = 5):
        """5 = "Title Only", 0 = "Title Slide" (standart shablon maketlari)."""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[layout_idx])
        self.page += 1
        add_transition(slide)
        return slide

    def _content_slide(self, title: str):
        slide = self._new(5)
        set_solid_bg(slide, WHITE)
        kicker = self.topic.upper()
        if len(kicker) > 70:
            kicker = kicker[:67].rstrip() + "..."
        add_text(slide, MX, Inches(0.32), CW - Inches(1.2), Inches(0.3), [kicker], size=11, color=self.pal["primary"],
                 bold=True, name="Kicker")
        set_title(slide, title, MX, TITLE_Y, CW, TITLE_H, (34, 32, 30, 28, 26, 24), self.pal["dark"], self.head,
                  anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, SW - MX - Inches(1.0), Inches(6.98), Inches(1.0), Inches(0.3), [str(self.page)], size=11,
                 color=MUTED, align=PP_ALIGN.RIGHT, name="Slide Number")
        return slide

    # -------- title / agenda / end --------
    def title_slide(self):
        pal = self.pal
        slide = self._new(0)
        set_gradient_bg(slide, pal["dark"], pal["primary"], 35)
        add_box(slide, SW - Inches(4.6), Inches(-1.9), Inches(7.0), Inches(7.0), WHITE, MSO_SHAPE.OVAL, alpha=7,
                name="Decor 1")
        add_box(slide, SW - Inches(2.9), Inches(3.9), Inches(4.4), Inches(4.4), pal["accent"], MSO_SHAPE.OVAL,
                alpha=22, name="Decor 2")

        pill = add_box(slide, Inches(0.85), Inches(1.45), Inches(2.3), Inches(0.42), pal["accent"],
                       MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5, name="Label")
        _prep_frame(pill.text_frame, MSO_ANCHOR.MIDDLE)
        p = pill.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = "TAQDIMOT"
        _style_run(r, 12, pal["dark"], bold=True)

        set_title(slide, self.topic, Inches(0.85), Inches(2.1), Inches(8.0), Inches(2.5), (54, 48, 44, 40, 36, 32),
                  WHITE, self.head, anchor=MSO_ANCHOR.BOTTOM)
        subtitle = self.deck.get("subtitle") or ""
        sub_ph = slide.placeholders[1]
        if not subtitle:
            sub_ph._element.getparent().remove(sub_ph._element)
        else:
            self._fill_subtitle(sub_ph, subtitle)
        add_text(slide, Inches(0.85), Inches(6.55), Inches(5), Inches(0.35), [str(date.today().year)], size=13,
                 color=mix(WHITE, pal["primary"], 0.35), name="Year")
        set_notes(slide, f"Assalomu alaykum! Bugungi taqdimotimiz mavzusi: {self.topic}. {subtitle}".strip())

    def _fill_subtitle(self, sub_ph, subtitle: str):
        sub_ph.left, sub_ph.top, sub_ph.width, sub_ph.height = Inches(0.85), Inches(4.85), Inches(7.6), Inches(1.1)
        tf = sub_ph.text_frame
        _prep_frame(tf)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        _no_bullet(p)
        r = p.add_run()
        r.text = subtitle
        _style_run(r, fit_size([subtitle], Inches(7.6), Inches(1.1), (22, 20, 18, 16)),
                   mix(self.pal["tint"], self.pal["accent"], 0.15))

    def agenda_slide(self, items: List[str]):
        pal = self.pal
        slide = self._content_slide("Reja")
        n = len(items)
        cols = 1 if n <= 5 else 2
        rows = math.ceil(n / cols)
        col_w = (CW - GAP * (cols - 1)) // cols
        row_h = min(Inches(0.92), CH // rows)
        d = min(Inches(0.56), row_h - Inches(0.16))
        text_w = col_w - d - Inches(0.25)
        size = min(fit_size([it], text_w, row_h, (22, 20, 18, 17, 16, 15, 14), line_spacing=1.0) for it in items)
        for i, item in enumerate(items):
            col, row = divmod(i, rows)
            x = MX + col * (col_w + GAP)
            y = CY + row * row_h
            add_box(slide, x, y + Inches(0.06), col_w, row_h - Inches(0.12), pal["tint"], MSO_SHAPE.ROUNDED_RECTANGLE,
                    radius=0.18, name=f"Agenda row {i + 1}")
            add_number_circle(slide, x + Inches(0.12), y + (row_h - d) // 2, d, i + 1, pal["primary"])
            add_text(slide, x + Inches(0.12) + d + Inches(0.2), y + Inches(0.06), text_w - Inches(0.2),
                     row_h - Inches(0.12), [item], size=size, color=TEXT, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0,
                     name=f"Agenda item {i + 1}")
        set_notes(slide, "Taqdimot quyidagi bo'limlardan iborat: " + "; ".join(items) + ".")

    def end_slide(self):
        pal = self.pal
        slide = self._new(5)
        set_gradient_bg(slide, pal["primary"], pal["dark"], 215)
        add_box(slide, Inches(-1.6), Inches(-1.6), Inches(4.6), Inches(4.6), pal["accent"], MSO_SHAPE.OVAL, alpha=22,
                name="Decor 1")
        add_box(slide, SW - Inches(3.4), SH - Inches(3.4), Inches(5.6), Inches(5.6), WHITE, MSO_SHAPE.OVAL, alpha=7,
                name="Decor 2")
        set_title(slide, "E'tiboringiz uchun rahmat!", Inches(1), Inches(2.35), SW - Inches(2), Inches(1.4),
                  (48, 44, 40), WHITE, self.head, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)
        add_text(slide, Inches(1), Inches(3.95), SW - Inches(2), Inches(0.6), ["Savollaringiz bo'lsa, marhamat"],
                 size=22, color=mix(pal["tint"], pal["accent"], 0.2), align=PP_ALIGN.CENTER, name="Questions")
        add_text(slide, Inches(1), Inches(6.45), SW - Inches(2), Inches(0.4), [self.topic], size=13,
                 color=mix(WHITE, pal["primary"], 0.35), align=PP_ALIGN.CENTER, name="Topic")
        set_notes(slide, "E'tiboringiz uchun rahmat! Endi savollaringizga javob berishga tayyorman.")

    def image_slide(self, path: str, caption: str):
        pal = self.pal
        slide = self._new(5)
        set_solid_bg(slide, pal["dark"])
        add_picture_cover(slide, path, 0, 0, SW, SH, name="Full image")
        add_box(slide, 0, SH - Inches(1.5), SW, Inches(1.5), pal["dark"], alpha=70, name="Caption backdrop")
        set_title(slide, caption, MX, SH - Inches(1.3), CW, Inches(1.0), (30, 28, 26, 24), WHITE, self.head,
                  anchor=MSO_ANCHOR.MIDDLE)
        # Title placeholder rasm ostida qolmasligi uchun uni eng yuqori qatlamga ko'chiramiz.
        title_el = slide.shapes.title._element
        tree = title_el.getparent()
        tree.remove(title_el)
        tree.append(title_el)

    # -------- content layouts --------
    def bullets_slide(self, s: dict, image: Optional[str] = None, check: bool = False):
        pal = self.pal
        slide = self._content_slide(s["title"])
        bullets = s["bullets"]
        highlight = s.get("highlight") or ""
        char = "✓" if check else "•"
        if image:
            text_w = Inches(7.0)
            img_x = MX + text_w + Inches(0.45)
            add_picture_cover(slide, image, img_x, CY, SW - MX - img_x, CH, rounded=True)
            hl_h = Inches(1.25) if highlight else 0
            max_h = CH - hl_h - (GAP if hl_h else 0)
            size = fit_size(bullets, text_w, max_h, BODY_SIZES, 1.1, 0.6, indent=Pt(BODY_SIZES[0] * 1.1))
            # Matn ustuni rasm balandligi bo'yicha vertikal markazlashtiriladi.
            used = min(max_h, int(text_height(bullets, text_w, size, 1.1, 0.6, indent=Pt(size * 1.1)) / 0.9))
            y0 = CY + (CH - used - (GAP + hl_h if hl_h else 0)) // 2
            add_bullets(slide, MX, y0, text_w, used, bullets, pal, sizes=(size,), char=char)
            if highlight:
                self._highlight_card(slide, MX, y0 + used + GAP, text_w, hl_h, highlight, compact=True)
        elif highlight:
            text_w = Inches(7.45)
            add_bullets(slide, MX, CY, text_w, CH, bullets, pal, char=char)
            card_x = MX + text_w + Inches(0.45)
            self._highlight_card(slide, card_x, CY, SW - MX - card_x, CH, highlight)
        else:
            self._numbered_rows(slide, bullets, check)
        return slide

    def _numbered_rows(self, slide, bullets: List[str], check: bool = False):
        """Rasm ham, ajratilgan fikr ham bo'lmagan slayd matn-devor bo'lib qolmasligi uchun: har bir punkt
        raqamli (yoki belgili) doira bilan alohida kartochka-qatorda."""
        pal = self.pal
        n = len(bullets)
        gap = Inches(0.18)
        row_h = min(Inches(1.15), (CH - gap * (n - 1)) // n)
        d = min(Inches(0.62), row_h - Inches(0.3))
        pad = Inches(0.25)
        text_x_off = pad + d + Inches(0.3)
        text_w = CW - text_x_off - pad
        size = min(fit_size([b], text_w, row_h - Inches(0.16), (24, 22, 20, 19, 18, 17, 16, 15, 14),
                            line_spacing=1.05) for b in bullets)
        block_h = n * row_h + (n - 1) * gap
        y0 = CY + (CH - block_h) // 2
        for i, text in enumerate(bullets):
            y = y0 + i * (row_h + gap)
            add_box(slide, MX, y, CW, row_h, pal["tint"], MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.14,
                    name=f"Row {i + 1}")
            add_number_circle(slide, MX + pad, y + (row_h - d) // 2, d, "\u2713" if check else i + 1,
                              pal["primary"])
            box = slide.shapes.add_textbox(MX + text_x_off, y + Inches(0.08), text_w, row_h - Inches(0.16))
            box.name = f"Row {i + 1} text"
            _prep_frame(box.text_frame, MSO_ANCHOR.MIDDLE, autofit=True)
            _write_bullets(box.text_frame, [text], size, TEXT, pal["primary"], lead_color=pal["dark"])
            _strip_bullet(box.text_frame.paragraphs[0])

    def _highlight_card(self, slide, x, y, w, h, text: str, compact: bool = False):
        pal = self.pal
        pad = Inches(0.3)
        if not compact:
            size = fit_size([text], w - 2 * pad, h - Inches(1.35) - pad, (24, 22, 20, 19, 18, 17, 16, 15, 14),
                            line_spacing=1.1)
            needed = text_height([text], w - 2 * pad, size, 1.1)
            h = min(h, Inches(1.35) + needed + pad + Inches(0.25))
        add_box(slide, x, y, w, h, pal["tint"], MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, name="Highlight card")
        if compact:
            add_text(slide, x + pad, y, Inches(0.6), h, ["“"], size=54, color=pal["primary"], font="Georgia",
                     bold=True, anchor=MSO_ANCHOR.MIDDLE, name="Quote mark")
            tx, tw, ty, th = x + pad + Inches(0.6), w - 2 * pad - Inches(0.6), y + Inches(0.12), h - Inches(0.24)
            size = fit_size([text], tw, th, (18, 17, 16, 15, 14), line_spacing=1.05)
            add_text(slide, tx, ty, tw, th, [text], size=size, color=pal["dark"], italic=True, bold=True,
                     anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05, autofit=True, name="Highlight")
            return
        add_text(slide, x + pad, y + Inches(0.2), Inches(1.2), Inches(1.1), ["“"], size=88, color=pal["primary"],
                 font="Georgia", bold=True, name="Quote mark")
        tx, tw = x + pad, w - 2 * pad
        ty, th = y + Inches(1.35), h - Inches(1.35) - pad
        size = fit_size([text], tw, th, (24, 22, 20, 19, 18, 17, 16, 15, 14), line_spacing=1.1)
        add_text(slide, tx, ty, tw, th, [text], size=size, color=pal["dark"], italic=True, bold=True,
                 line_spacing=1.1, autofit=True, name="Highlight")

    def cards_slide(self, s: dict):
        pal = self.pal
        slide = self._content_slide(s["title"])
        items = s["items"]
        n = len(items)
        cols = n if n <= 4 else 3
        rows = math.ceil(n / cols)
        card_w = (CW - GAP * (cols - 1)) // cols
        max_card_h = (CH - GAP * (rows - 1)) // rows
        pad = Inches(0.28)
        d = Inches(0.55)
        inner_w = card_w - 2 * pad
        head_w = inner_w - d - Inches(0.2)
        head_size = min(fit_size([it["title"]], head_w, Inches(0.8), (22, 20, 19, 18, 17, 16, 15), line_spacing=1.0,
                                 char_w=0.55) for it in items)
        max_text_h = max_card_h - pad * 2 - Inches(0.95)
        text_size = min(fit_size([it["text"]], inner_w, max_text_h, (20, 19, 18, 17, 16, 15, 14), line_spacing=1.1)
                        for it in items)
        needed = max(text_height([it["text"]], inner_w, text_size, 1.1) for it in items)
        text_h = min(max_text_h, int(needed / 0.9) + Inches(0.1))
        card_h = pad * 2 + Inches(0.95) + text_h
        y0 = CY + (CH - (rows * card_h + (rows - 1) * GAP)) // 2
        for i, it in enumerate(items):
            r, c = divmod(i, cols)
            x = MX + c * (card_w + GAP)
            y = y0 + r * (card_h + GAP)
            add_box(slide, x, y, card_w, card_h, pal["tint"], MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06,
                    name=f"Card {i + 1}")
            add_number_circle(slide, x + pad, y + pad + (Inches(0.8) - d) // 2, d, i + 1, pal["primary"])
            add_text(slide, x + pad + d + Inches(0.2), y + pad, head_w, Inches(0.8), [it["title"]], size=head_size,
                     color=pal["dark"], bold=True, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0,
                     name=f"Card {i + 1} title")
            add_text(slide, x + pad, y + pad + Inches(0.95), inner_w, text_h, [it["text"]], size=text_size,
                     color=TEXT, autofit=True, name=f"Card {i + 1} text")
        return slide

    def stats_slide(self, s: dict):
        pal = self.pal
        slide = self._content_slide(s["title"])
        items = s["items"]
        text = s.get("text") or ""
        n = len(items)
        card_w = (CW - GAP * (n - 1)) // n
        card_h = Inches(2.7) if text else Inches(3.6)
        pad = Inches(0.3)
        text_size = fit_size([text], CW, CH - card_h - Inches(0.4), (20, 19, 18, 17, 16, 15, 14), line_spacing=1.15)
        text_block = int(text_height([text], CW, text_size, 1.15) / 0.9) + Inches(0.4) if text else 0
        y = CY + (CH - card_h - text_block) // 2
        inner_w = card_w - 2 * pad
        val_size = min(fit_single_line(it["value"], inner_w, (60, 54, 48, 44, 40, 36, 32, 28)) for it in items)
        val_h = Inches(val_size * 1.3 / 72)
        label_h = card_h - pad * 2 - val_h - Inches(0.1)
        label_size = min(fit_size([it["label"]], inner_w, label_h, (17, 16, 15, 14, 13), line_spacing=1.05)
                         for it in items)
        for i, it in enumerate(items):
            x = MX + i * (card_w + GAP)
            add_box(slide, x, y, card_w, card_h, pal["tint"], MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06,
                    name=f"Stat {i + 1}")
            add_text(slide, x + pad, y + pad, inner_w, val_h, [it["value"]], size=val_size, color=pal["primary"],
                     font=self.head, bold=True, anchor=MSO_ANCHOR.BOTTOM, line_spacing=1.0,
                     name=f"Stat {i + 1} value")
            add_text(slide, x + pad, y + pad + val_h + Inches(0.1), inner_w, label_h, [it["label"]], size=label_size,
                     color=TEXT, line_spacing=1.05, autofit=True, name=f"Stat {i + 1} label")
        if text:
            ty = y + card_h + Inches(0.4)
            add_text(slide, MX, ty, CW, CB - ty, [text], size=text_size, color=TEXT, line_spacing=1.15, autofit=True,
                     name="Stat commentary")
        return slide

    def timeline_slide(self, s: dict):
        pal = self.pal
        slide = self._content_slide(s["title"])
        items = s["items"]
        n = len(items)
        col_w = (CW - GAP * (n - 1)) // n
        label_h, dot = Inches(0.6), Inches(0.32)
        head_h = label_h + Inches(0.3) + dot + Inches(0.35)
        max_text_h = CH - head_h
        label_size = min(fit_single_line(it["label"], col_w, (28, 26, 24, 22, 20, 18, 16)) for it in items)
        text_size = min(fit_size([it["text"]], col_w - Inches(0.1), max_text_h, (20, 19, 18, 17, 16, 15, 14),
                                 line_spacing=1.1) for it in items)
        needed = max(text_height([it["text"]], col_w - Inches(0.1), text_size, 1.1) for it in items)
        text_h = min(max_text_h, int(needed / 0.9) + Inches(0.1))
        label_y = CY + (CH - head_h - text_h) // 2
        line_y = label_y + label_h + Inches(0.3)
        text_y = line_y + dot + Inches(0.35)
        add_box(slide, MX, line_y + dot // 2 - Pt(1.5), CW, Pt(3), mix(pal["primary"], WHITE, 0.7),
                name="Timeline axis")
        for i, it in enumerate(items):
            x = MX + i * (col_w + GAP)
            add_text(slide, x, label_y, col_w, label_h, [it["label"]], size=label_size, color=pal["primary"],
                     font=self.head, bold=True, anchor=MSO_ANCHOR.BOTTOM, line_spacing=1.0, name=f"Step {i + 1} label")
            add_box(slide, x, line_y, dot, dot, pal["primary"], MSO_SHAPE.OVAL, name=f"Step {i + 1} dot")
            add_box(slide, x + Inches(0.09), line_y + Inches(0.09), dot - Inches(0.18), dot - Inches(0.18), WHITE,
                    MSO_SHAPE.OVAL, name=f"Step {i + 1} dot center")
            add_text(slide, x, text_y, col_w - Inches(0.1), text_h, [it["text"]], size=text_size, color=TEXT,
                     autofit=True, name=f"Step {i + 1} text")
        return slide

    def compare_slide(self, s: dict):
        pal = self.pal
        slide = self._content_slide(s["title"])
        col_w = (CW - Inches(0.4)) // 2
        head_h = Inches(0.75)
        pad = Inches(0.3)
        max_body_h = CH - head_h - Inches(0.12)
        sides = [(s["left"], pal["primary"]), (s["right"], pal["dark"])]
        size = min(fit_size(side["points"], col_w - 2 * pad, max_body_h - 2 * pad, (22, 20, 19, 18, 17, 16, 15, 14),
                            line_spacing=1.1, space_after_ratio=0.6, indent=Pt(24)) for side, _ in sides)
        needed = max(text_height(side["points"], col_w - 2 * pad, size, 1.1, 0.6, indent=Pt(size * 1.1))
                     for side, _ in sides)
        body_h = min(max_body_h, int(needed / 0.9) + 2 * pad + Inches(0.1))
        top = CY + (CH - head_h - Inches(0.12) - body_h) // 2
        body_y = top + head_h + Inches(0.12)
        for i, (side, color) in enumerate(sides):
            x = MX + i * (col_w + Inches(0.4))
            head = add_box(slide, x, top, col_w, head_h, color, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18,
                           name=f"Column {i + 1} header")
            _prep_frame(head.text_frame, MSO_ANCHOR.MIDDLE, margin=pad)
            p = head.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = side["title"]
            _style_run(r, fit_single_line(side["title"], col_w - 2 * pad, (22, 20, 18, 16)), WHITE, self.head, bold=True)
            add_box(slide, x, body_y, col_w, body_h, pal["tint"], MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05,
                    name=f"Column {i + 1} body")
            box = slide.shapes.add_textbox(x + pad, body_y + pad, col_w - 2 * pad, body_h - 2 * pad)
            box.name = f"Column {i + 1} points"
            _prep_frame(box.text_frame, autofit=True)
            _write_bullets(box.text_frame, side["points"], size, TEXT, color, lead_color=pal["dark"])
        return slide

    def quote_slide(self, s: dict):
        pal = self.pal
        slide = self._new(5)
        set_solid_bg(slide, pal["tint"])
        add_text(slide, MX, Inches(0.7), Inches(2), Inches(1.6), ["“"], size=140, color=pal["primary"],
                 font="Georgia", bold=True, name="Quote mark")
        text = s["text"]
        set_title(slide, text, Inches(1.4), Inches(2.0), SW - Inches(2.8), Inches(3.2), (36, 32, 30, 28, 26, 24),
                  pal["dark"], self.head, anchor=MSO_ANCHOR.MIDDLE)
        slide.shapes.title.text_frame.paragraphs[0].runs[0].font.italic = True
        if s.get("author"):
            add_text(slide, Inches(1.4), Inches(5.45), SW - Inches(2.8), Inches(0.5), ["— " + s["author"]],
                     size=20, color=pal["primary"], bold=True, name="Quote author")
        add_text(slide, SW - MX - Inches(1.0), Inches(6.98), Inches(1.0), Inches(0.3), [str(self.page)], size=11,
                 color=MUTED, align=PP_ALIGN.RIGHT, name="Slide Number")
        return slide

    # -------- assemble --------
    def build(self, slide_images: List[str], large_images: List[str]) -> Presentation:
        slides = self.deck["slides"]
        self.title_slide()
        agenda = [s["title"] for s in slides if s["type"] != "quote"]
        self.agenda_slide(agenda)

        # "Matn" rasmlari punktli slaydlarga, ortib qolganlari katta rasm slaydlariga ketadi.
        text_imgs = list(slide_images)
        image_for = {}
        for i, s in enumerate(slides):
            if text_imgs and s["type"] in ("bullets", "conclusion"):
                image_for[i] = text_imgs.pop(0)
        big_imgs = list(large_images) + text_imgs
        # Katta rasmlar kontent orasiga teng taqsimlanadi.
        big_at = {}
        for k, path in enumerate(big_imgs):
            pos = round((k + 1) * len(slides) / (len(big_imgs) + 1))
            big_at.setdefault(pos, []).append(path)

        renderers = {
            "bullets": lambda s, i: self.bullets_slide(s, image_for.get(i)),
            "conclusion": lambda s, i: self.bullets_slide(s, image_for.get(i), check=True),
            "cards": lambda s, i: self.cards_slide(s),
            "stats": lambda s, i: self.stats_slide(s),
            "timeline": lambda s, i: self.timeline_slide(s),
            "compare": lambda s, i: self.compare_slide(s),
            "quote": lambda s, i: self.quote_slide(s),
        }
        for i, s in enumerate(slides):
            for path in big_at.get(i, []):
                try:
                    self.image_slide(path, s.get("title") or self.topic)
                except Exception:
                    logging.exception("Katta rasm slaydini qo'shib bo'lmadi")
            try:
                slide = renderers[s["type"]](s, i)
            except Exception:
                logging.exception("Slayd (%s) chizishda xatolik, oddiy ko'rinishga o'tildi", s.get("type"))
                slide = self.bullets_slide({"title": s.get("title") or self.topic,
                                            "bullets": _fallback_bullets(s)})
            set_notes(slide, s.get("notes", ""))
        for path in big_at.get(len(slides), []):
            try:
                self.image_slide(path, self.topic)
            except Exception:
                logging.exception("Katta rasm slaydini qo'shib bo'lmadi")
        self.end_slide()

        cp = self.prs.core_properties
        cp.title = self.topic
        cp.subject = self.deck.get("subtitle") or self.topic
        cp.author = "PPT Yordamchi"
        cp.language = LANG
        return self.prs


def _fallback_bullets(s: dict) -> List[str]:
    out = list(s.get("bullets") or [])
    for it in s.get("items") or []:
        out.append(" ".join(v for v in (it.get("title") or it.get("label") or it.get("value"),
                                         it.get("text") or it.get("label")) if v))
    if s.get("text"):
        out.append(s["text"])
    return out[:6] or ["—"]


def safe_filename(topic: str) -> str:
    name = re.sub(r"[^\w\s\-]", "", topic, flags=re.UNICODE).strip()
    name = re.sub(r"\s+", "_", name)[:60]
    return (name or "taqdimot") + ".pptx"


def create_presentation_file(deck: dict, out_dir: str, slide_images: List[str],
                             large_images: List[str]) -> Tuple[str, int]:
    """Faylni saqlaydi va (yo'l, slaydlar soni) qaytaradi."""
    prs = DeckBuilder(deck).build(slide_images, large_images)
    path = os.path.join(out_dir, safe_filename(deck["title"]))
    prs.save(path)
    return path, len(prs.slides)
